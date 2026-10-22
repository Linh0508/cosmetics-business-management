from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"c:\Users\linh0\OneDrive\Desktop\DATN\cosmetics-business-management")
EXCEL_PATH = ROOT / "data" / "Data mỹ phẩm.xlsx"
IMAGE_DIR = ROOT / "data" / "Hình ảnh mỹ phẩm"
CSV_PATH = ROOT / "data" / "products_clean.csv"
JSON_PATH = ROOT / "data" / "products_clean.json"


def normalize_text(value):
    if value is None:
        return ""
    text = str(value).strip()
    text = text.replace("\xa0", " ")
    text = text.replace("\r", " ").replace("\n", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_key(value):
    text = normalize_text(value).lower()
    text = text.replace("đ", "d")
    text = text.replace("&", " and ")
    text = text.replace("'", "")
    text = re.sub(r"[^a-z0-9\s_\-]+", "", text)
    text = re.sub(r"\s+", "", text)
    return text.strip()


def match_key(key, candidate):
    a = normalize_key(key)
    b = normalize_key(candidate)
    if not a or not b:
        return False
    return a in b or b in a or a == b


def get_value_by_candidates(row, *candidates):
    for candidate in candidates:
        for key, value in row.items():
            if normalize_key(key) == normalize_key(candidate):
                return value
    for candidate in candidates:
        for key, value in row.items():
            if normalize_key(candidate) in normalize_key(key) or normalize_key(key) in normalize_key(candidate):
                return value
    return None


def split_by_variants(value):
    if value is None:
        return []
    text = normalize_text(value)
    if not text:
        return []
    text = text.replace(";", "_").replace("|", "_").replace("/", "_")
    text = text.replace(" - ", "_").replace("-", "_")
    text = text.replace(" _ ", "_").replace("_ ", "_").replace(" _", "_")
    text = re.sub(r"\s+", " ", text)
    parts = [part.strip() for part in text.split("_") if part.strip()]
    return parts


def normalize_capacity_variant(raw):
    if raw is None:
        return ""
    text = normalize_text(raw)
    if not text:
        return ""
    text = text.lower()
    text = text.replace(" ", "")
    text = re.sub(r"(?<=\d)(ml|l|g|mg|kg)", r"\1", text, flags=re.I)
    text = re.sub(r"[^0-9a-z%]+", "", text)
    return text.strip()


def clean_price_number(raw):
    if raw is None:
        return ""
    text = str(raw).strip()
    if not text:
        return ""

    text = text.lower().replace("vnd", " ").replace("₫", " ").replace("đ", " ")
    text = re.sub(r"[^0-9,\._/|;\-\s]+", " ", text)
    text = text.replace("/", "_").replace("|", "_").replace(";", "_").replace("-", "_")
    text = re.sub(r"\s+", " ", text).strip()

    variant_parts = []
    for part in text.split("_"):
        item = part.strip()
        if not item:
            continue
        digits = re.findall(r"\d+", item)
        if not digits:
            continue
        variant_parts.append("".join(digits))

    if not variant_parts:
        return ""
    return "_".join(variant_parts)


def parse_price_capacity_pairs(raw_price, raw_capacity):
    price_parts = [clean_price_number(part) for part in split_by_variants(raw_price) if clean_price_number(part)]
    capacity_parts = [normalize_capacity_variant(part) for part in split_by_variants(raw_capacity) if normalize_capacity_variant(part)]

    if not price_parts and not capacity_parts:
        return []

    if not capacity_parts:
        capacity_parts = ["macdinh"]

    if not price_parts:
        price_parts = ["0"]

    max_count = max(len(capacity_parts), len(price_parts))
    variants = []
    for idx in range(max_count):
        capacity = capacity_parts[idx].strip() if idx < len(capacity_parts) else (capacity_parts[0] if len(capacity_parts) == 1 else f"phiensanpham{idx + 1}")
        price = price_parts[idx].strip() if idx < len(price_parts) else (price_parts[0] if len(price_parts) == 1 else "")
        if not capacity and not price:
            continue
        variants.append({
            "capacity": capacity,
            "price": price,
            "raw_price": price,
        })

    if not variants:
        variants.append({"capacity": "macdinh", "price": clean_price_number(raw_price), "raw_price": clean_price_number(raw_price)})

    return variants


def normalize_key_for_match(value):
    return re.sub(r"[^a-z0-9]+", "", normalize_text(value).lower())


def matches_variant_name(candidate_key: str, stem_key: str) -> bool:
    if not candidate_key or not stem_key:
        return False
    if stem_key == candidate_key:
        return True
    for prefix in [candidate_key + "_", candidate_key + " "]:
        if stem_key.startswith(prefix):
            suffix = stem_key[len(prefix):]
            if suffix == "" or suffix.isdigit():
                return True
    return False


def find_image_paths(product_name: str, raw_image_field: object) -> list[str]:
    candidates: list[str] = []
    if product_name:
        candidates.append(str(product_name))
    if raw_image_field is not None:
        raw = str(raw_image_field)
        if raw:
            candidates.append(raw)

    if not candidates:
        return []

    image_files = []
    for candidate in candidates:
        candidate_key = normalize_key_for_match(candidate)
        if not candidate_key:
            continue
        for p in IMAGE_DIR.iterdir():
            if not p.is_file():
                continue
            stem_key = normalize_key_for_match(p.stem)
            if not stem_key:
                continue
            if (
                candidate_key == stem_key
                or candidate_key in stem_key
                or stem_key in candidate_key
                or matches_variant_name(candidate_key, stem_key)
            ):
                image_files.append(p.relative_to(ROOT).as_posix())

    if not image_files and product_name:
        product_key = normalize_key_for_match(product_name)
        for p in IMAGE_DIR.iterdir():
            if not p.is_file():
                continue
            stem_key = normalize_key_for_match(p.stem)
            if not stem_key:
                continue
            if matches_variant_name(product_key, stem_key):
                image_files.append(p.relative_to(ROOT).as_posix())

    unique = []
    seen = set()
    for item in image_files:
        if item not in seen:
            seen.add(item)
            unique.append(item)
    return sorted(unique)


def load_rows_from_excel():
    try:
        wb = load_workbook(EXCEL_PATH, data_only=True)
    except PermissionError:
        return []

    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = []
    for idx, value in enumerate(rows[0]):
        headers.append("col_{}".format(idx) if value is None else str(value).strip())

    records = []
    for row in rows[1:]:
        if row is None or all(value is None or str(value).strip() == "" for value in row):
            continue
        record = {}
        for idx, header in enumerate(headers):
            value = row[idx] if idx < len(row) else None
            record[header] = value
        records.append(record)
    return records


def load_rows_from_csv():
    if not CSV_PATH.exists():
        return []
    with CSV_PATH.open("r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        return list(reader)


def build_record(row):
    product_name = get_value_by_candidates(row, "Sản phẩm", "Tên sản phẩm", "product_name") or ""
    category = get_value_by_candidates(row, "Danh mục", "category")
    brand = get_value_by_candidates(row, "Thương hiệu", "brand")
    raw_price = get_value_by_candidates(row, "Giá bán", "price", "Giá")
    raw_capacity = get_value_by_candidates(row, "Dung tích", "capacity")
    unit = get_value_by_candidates(row, "Đơn vị", "unit")
    skin_type = get_value_by_candidates(row, "Loại da phù hợp", "skin_type", "Loại da")
    usage = get_value_by_candidates(row, "Công dụng", "Công dụng ", "usage")
    ingredients = get_value_by_candidates(row, "Thành phần", "ingredients")
    origin = get_value_by_candidates(row, "Xuất xứ", "origin")
    expiry = get_value_by_candidates(row, "Hạn sử dụng", "expiry")
    raw_image = get_value_by_candidates(row, "Dường dẫn ảnh", "image_paths")

    image_paths = find_image_paths(str(product_name), raw_image)
    variants = parse_price_capacity_pairs(raw_price, raw_capacity)
    normalized_prices = [variant["price"] for variant in variants if variant.get("price")]
    normalized_capacities = [variant["capacity"] for variant in variants if variant.get("capacity")]
    selected_price = "_".join(normalized_prices) if normalized_prices else clean_price_number(raw_price)
    selected_capacity = "_".join(normalized_capacities) if normalized_capacities else str(raw_capacity or "")

    return {
        "product_name": str(product_name).strip(),
        "category": str(category).strip() if category is not None else "",
        "brand": str(brand).strip() if brand is not None else "",
        "price": selected_price,
        "capacity": selected_capacity,
        "unit": str(unit).strip() if unit is not None else "",
        "skin_type": str(skin_type).strip() if skin_type is not None else "",
        "usage": str(usage).strip() if usage is not None else "",
        "ingredients": str(ingredients).strip() if ingredients is not None else "",
        "origin": str(origin).strip() if origin is not None else "",
        "expiry": str(expiry).strip() if expiry is not None else "",
        "image_paths": image_paths,
        "image_count": len(image_paths),
        "capacity_variants": variants,
    }


def main():
    rows = load_rows_from_excel() or load_rows_from_csv()
    cleaned = []
    missing = []

    for row in rows:
        record = build_record(row)
        if not record["product_name"]:
            continue

        if not record["image_paths"]:
            missing.append({
                "product": record["product_name"],
                "raw_image": row.get("Dường dẫn ảnh") if isinstance(row, dict) else "",
            })

        cleaned.append(record)

    for record in cleaned:
        record.pop("capacity_variants", None)

    with CSV_PATH.open("w", encoding="utf-8", newline="") as f:
        fieldnames = [
            "product_name",
            "category",
            "brand",
            "price",
            "capacity",
            "unit",
            "skin_type",
            "usage",
            "ingredients",
            "origin",
            "expiry",
            "image_count",
            "image_paths",
        ]
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for record in cleaned:
            row = {key: record.get(key, "") for key in fieldnames}
            row["image_paths"] = "; ".join(record.get("image_paths", []))
            writer.writerow(row)

    with JSON_PATH.open("w", encoding="utf-8") as f:
        json.dump(cleaned, f, ensure_ascii=False, indent=2)

    missing_report = ROOT / "data" / "missing_image_products.json"
    with missing_report.open("w", encoding="utf-8") as f:
        json.dump(missing, f, ensure_ascii=False, indent=2)

    print(f"Processed {len(cleaned)} products")
    print(f"Exported CSV: {CSV_PATH}")
    print(f"Exported JSON: {JSON_PATH}")
    print(f"Missing image mappings: {len(missing)}")
    print(f"Missing image report: {missing_report}")


if __name__ == "__main__":
    main()
