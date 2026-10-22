from pathlib import Path
from openpyxl import load_workbook

path = Path(r'c:\Users\linh0\OneDrive\Desktop\DATN\cosmetics-business-management\data\Data mỹ phẩm.xlsx')
print('exists=', path.exists())
try:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    print('sheets=', wb.sheetnames)
    print('max_row=', ws.max_row)
    print('max_col=', ws.max_column)
    first = next(ws.iter_rows(values_only=True, max_row=5), [])
    print('sample=', first)
except Exception as e:
    print('ERROR:', type(e).__name__, str(e))
    raise
